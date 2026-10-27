"""E/F: native-file inspection, actual PowerPoint rendering, final review gate."""
from pathlib import Path
import hashlib
import shutil
import uuid
import sys
from PIL import Image, ImageStat
from pptx import Presentation
from pypdf import PdfReader
from .common import ROOT, command, read, require, sha, write
from .contracts import validate_plan, validate_science

def current_design(job):
    validate_plan(job,validate_science(job))
    spec=read(job/'D/deck-spec.json')
    if read(job/'run.json').get('workflowContract')=='generalized-editorial-v1':
        require(spec.get('workflowContract')=='generalized-editorial-v1','WORKFLOW_CONTRACT','A new run must retain its generalized design/visual gates.')
    require(spec['planSha256']==sha(job/'C/plan.json') and spec['analysisSha256']==sha(job/'B/analysis.json')
        and spec['scientificReviewSha256']==sha(job/'B/review.json') and spec['narrativeReviewSha256']==sha(job/'C/review.json'),
        'STALE_DESIGN','Upstream analysis/plan/review changed; regenerate the design.')
    report=read(job/'D/design-report.json')
    require(report.get('specSha256')==sha(job/'D/deck-spec.json') and not report['issues'],'STALE_DESIGN','Design preflight is missing, stale or failed.')
    for slide in spec['slides']:
        for obj in slide['objects']:
            if obj['type']=='image':require(sha(obj['path'])==obj['sha256'],'ASSET_CHANGED','Original figure crop changed.')
    return spec

def build(job):
    spec=current_design(job)
    command(['node',ROOT/'scripts/build_deck.cjs',job/'D/deck-spec.json',job/'E/presentation.pptx'],job/'E/build.log')
    require((job/'E/presentation.pptx').is_file() and (job/'E/presentation.pptx').stat().st_size>1000,'EXPORT_FAILED','No valid PPTX was generated.')
    native_check(job,spec)
    write(job/'E/build-manifest.json',{'engine':'pptxgenjs','specSha256':sha(job/'D/deck-spec.json'),
        'pptxSha256':sha(job/'E/presentation.pptx')})

def native_check(job,spec):
    deck=Presentation(job/'E/presentation.pptx')
    require(len(deck.slides)==len(spec['slides']),'PPTX_SLIDES','Native PPTX slide count changed.')
    require(abs(deck.slide_width/deck.slide_height-16/9)<.001,'PPTX_ASPECT','PPTX must be 16:9.')
    report={'engine':'OOXML inspection','slides':len(deck.slides),'textShapes':0,'nativeCharts':0,'originalImages':0,'notes':0,'issues':[]}
    for index,(slide,expected) in enumerate(zip(deck.slides,spec['slides']),1):
        texts=[s.text for s in slide.shapes if s.has_text_frame and s.text]
        expected_text=[o['text'] for o in expected['objects'] if o['type']=='text' and o['text']]
        require(texts==expected_text,'PPTX_TEXT_FIDELITY',f'Native text differs from reviewed specification on slide {index}.')
        report['textShapes']+=len(texts)
        require(slide.has_notes_slide and expected['notes'] in slide.notes_slide.notes_text_frame.text,
            'PPTX_NOTES',f'Slide {index} speaker notes are missing/changed.')
        report['notes']+=1
        charts=[s.chart for s in slide.shapes if s.has_chart]
        expected_charts=[o for o in expected['objects'] if o['type']=='chart']
        require(len(charts)==len(expected_charts),'PPTX_CHART_EDITABILITY',f'Native chart lost on slide {index}.')
        for chart,target in zip(charts,expected_charts):
            require(list(chart.series[0].values)==target['values'] and [x.label for x in chart.plots[0].categories]==target['labels'],
                    'PPTX_CHART_FIDELITY',f'Native chart values/labels differ on slide {index}.')
            if spec.get('editorialQualityPolicy') and target['unit']=='%':
                bodies=chart._chartSpace.xpath('.//c:dLbls/c:txPr/a:bodyPr')
                require(bodies and all(b.get('wrap')=='none' for b in bodies),'PPTX_PERCENT_WRAP',
                    'Native percent chart labels require explicit no-wrap text bodies.')
        report['nativeCharts']+=len(charts)
        images=[s for s in slide.shapes if s.shape_type==13]
        expected_images=[o for o in expected['objects'] if o['type']=='image']
        require(sorted(hashlib.sha256(x.image.blob).hexdigest() for x in images)==sorted(x['sha256'] for x in expected_images),
            'PPTX_IMAGE_FIDELITY',f'Original source image changed or was lost on slide {index}.')
        report['originalImages']+=len(images)
        if spec.get('designSystem')=='modern-scientific-editorial':
            named={s.name:s for s in slide.shapes}
            for obj in expected['objects']:
                require(obj['id'] in named,'PPTX_OBJECT','Missing editable object '+obj['id'])
                shape=named[obj['id']]
                require(all(abs(a/914400-b)<.0001 for a,b in zip([shape.left,shape.top,shape.width,shape.height],[obj['x'],obj['y'],obj['w'],obj['h']])),
                    'PPTX_GEOMETRY',f'Native geometry changed: {obj["id"]}')
            for image in images:
                require(not any([image.crop_left,image.crop_right,image.crop_top,image.crop_bottom]),'PPTX_IMAGE_CROP','Original scientific figure was cropped.')
        for shape in slide.shapes:
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width+1000 and shape.top+shape.height<=deck.slide_height+1000,
                'PPTX_OVERFLOW',f'Object exceeds slide {index} canvas.')
    write(job/'E/native-check.json',report)
    return report

def qa(job):
    spec=current_design(job)
    build_manifest=read(job/'E/build-manifest.json')
    require(build_manifest['specSha256']==sha(job/'D/deck-spec.json') and build_manifest['pptxSha256']==sha(job/'E/presentation.pptx'),
            'STALE_BUILD','PPTX/spec changed after generation.')
    native=native_check(job,spec)
    render=job/'F/render'/uuid.uuid4().hex[:10]
    render.mkdir(parents=True)
    resolution=[2560,1440] if spec.get('designSystem')=='modern-scientific-editorial' else [1600,900]
    command([sys.executable,ROOT/'scripts/export_powerpoint.py',job/'E/presentation.pptx',
        '--output',render,'--width',resolution[0],'--height',resolution[1]],job/'F/powerpoint.log',timeout=300)
    pdf=render/'deck.pdf'
    require(pdf.is_file() and len(PdfReader(pdf).pages)==len(spec['slides']),'PDF_EXPORT','PDF export is absent or page count differs.')
    images=sorted((render/'slides').glob('slide-*.png'))
    require(len(images)==len(spec['slides']),'RENDER_MISSING','Some slides were not rendered.')
    pages=[]
    for slide,image in zip(spec['slides'],images):
        with Image.open(image) as im:
            require(list(im.size)==resolution,'RENDER_SIZE','Unexpected render size.')
            require(max(ImageStat.Stat(im.convert('RGB')).stddev)>2,'BLANK_RENDER',f'Blank image: {image.name}')
        pages.append({'slideId':slide['id'],'image':str(image.relative_to(job)).replace('\\','/'),'imageSha256':sha(image)})
    if spec.get('editorialQualityPolicy'):
        from .editorial_quality import postflight
        screen=postflight(job,spec,read(render/'export-result.json'),pdf)
        require(screen['status']!='blocked','EDITORIAL_RENDER_QUALITY',str(screen['findings']))
    report={'status':'awaiting-visual-review','pptxSha256':sha(job/'E/presentation.pptx'),
        'specSha256':sha(job/'D/deck-spec.json'),'pdf':str(pdf.relative_to(job)).replace('\\','/'),'pdfSha256':sha(pdf),
        'native':native,'slides':pages,'resolution':resolution,'warnings':['Timing is a plan, not measured delivery.',
        'Complex equations remain original image crops; native OMML is not implemented.',
        'HTML geometry and native shape checks do not establish visual/scientific correctness.'],
        'historicalFailuresPreserved':str(ROOT/'references/known-limitations.json')}
    if spec.get('editorialQualityPolicy'):report['editorialQualitySha256']=sha(job/'F/editorial-quality.json')
    write(job/'F/qa-report.json',report)
    template={'reviewer':'','qaSha256':sha(job/'F/qa-report.json'),
        'pptxSha256':report['pptxSha256'],'slides':[{'id':x['slideId'],'imageSha256':x['imageSha256'],
        'status':'pending','checks':[],'notes':''} for x in pages]}
    if spec.get('editorialQualityPolicy'):
        template.update(editorialQualitySha256=sha(job/'F/editorial-quality.json'),editorialChecks=[],editorialDispositions=[])
    if spec.get('workflowContract')=='generalized-editorial-v1':
        from .general_quality import REVIEW_CHECKS
        template.update(requiredEditorialChecks=sorted(REVIEW_CHECKS))
        for item,pair in zip(template['slides'],screen['generalizedScreens']['fontsAndPreviewPairs']):item['htmlPreviewSha256']=pair['htmlPreviewSha256']
    write(job/'F/visual-review.template.json',template)
    return report

def finalize(job):
    spec=current_design(job)
    qa_report=read(job/'F/qa-report.json')
    review=read(job/'F/visual-review.json')
    if spec.get('editorialQualityPolicy'):
        screen=read(job/'F/editorial-quality.json')
        require(qa_report.get('editorialQualitySha256')==sha(job/'F/editorial-quality.json') and
            screen['pptxSha256']==sha(job/'E/presentation.pptx') and screen['specSha256']==sha(job/'D/deck-spec.json') and
            screen['status']=='passed-automatic-screen-awaiting-visual-review','STALE_EDITORIAL_SCREEN','Missing, failed or stale actual-render editorial screen.')
        require(review.get('editorialQualitySha256')==sha(job/'F/editorial-quality.json') and
            set(review.get('editorialChecks',[]))>={'percentage-layout','figure-legibility','concise-sources','delivery-timing'},
            'EDITORIAL_REVIEW','Review all editorial screens against actual images and notes.')
        if spec.get('workflowContract')=='generalized-editorial-v1':
            from .general_quality import REVIEW_CHECKS
            require(set(review.get('editorialChecks',[]))>=REVIEW_CHECKS,'GENERAL_VISUAL_REVIEW','Review units, density/rhythm, fonts and paired HTML/PPTX images too.')
        decisions={(x.get('code'),x.get('slideId')):x for x in review.get('editorialDispositions',[])}
        for finding in [*screen['preflight']['findings'],*screen.get('findings',[])]:
            decision=decisions.get((finding['code'],finding['slideId']),{})
            require(decision.get('status')=='accepted-after-actual-review' and decision.get('rationale'),
                'EDITORIAL_REVIEW','Document an actual review decision for '+finding['code']+' on '+finding['slideId'])
    require(qa_report['pptxSha256']==sha(job/'E/presentation.pptx') and qa_report['specSha256']==sha(job/'D/deck-spec.json'),
        'STALE_QA','Generated file/spec changed; rerun actual rendering and QA.')
    require(review.get('qaSha256')==sha(job/'F/qa-report.json') and review.get('pptxSha256')==qa_report['pptxSha256'] and review.get('reviewer'),
        'STALE_VISUAL_REVIEW','Visual review is absent or stale.')
    items={s['id']:s for s in review.get('slides',[])}
    require(len(items)==len(spec['slides']),'INCOMPLETE_VISUAL_REVIEW','Review every actual PowerPoint render.')
    checks={'layout','legibility','scientific-figures','sources','notes'}
    if spec.get('workflowContract')=='generalized-editorial-v1':
        from .general_quality import PAGE_CHECKS
        checks=PAGE_CHECKS
        for pair in screen['generalizedScreens']['fontsAndPreviewPairs']:
            require(sha(job/pair['htmlPreview'])==pair['htmlPreviewSha256'] and items.get(pair['slideId'],{}).get('htmlPreviewSha256')==pair['htmlPreviewSha256'],
                'HTML_REVIEW_CHANGED','Review every current HTML preview alongside the actual PowerPoint PNG; the screenshot hash changed or was not reviewed.')
    for source in qa_report['slides']:
        item=items.get(source['slideId'],{})
        require(sha(job/source['image'])==source['imageSha256'] and item.get('imageSha256')==source['imageSha256'],
            'RENDER_CHANGED','Rendered image changed or has not been reviewed.')
        require(item.get('status')=='passed' and set(item.get('checks',[]))>=checks and item.get('notes'),
            'VISUAL_BLOCK',f'{source["slideId"]}: visual blockers or missing checks.')
    require(sha(job/qa_report['pdf'])==qa_report['pdfSha256'],'PDF_CHANGED','PDF changed after export.')
    # Publish locally only after every gate; incomplete attempts stay in A-F.
    destination=job/'deliverables'
    require(not destination.exists(),'ALREADY_FINALIZED','Final output is immutable; create a new run for a revision.')
    destination.mkdir()
    for source,target in [(job/'E/presentation.pptx','presentation.pptx'),(job/qa_report['pdf'],'presentation.pdf'),
        (job/'C/speaker-notes.md','speaker-notes.md'),(job/'B/evidence-report.md','evidence-report.md')]:
        shutil.copyfile(source,destination/target)
    (destination/'previews').mkdir()
    for i,item in enumerate(qa_report['slides'],1):shutil.copyfile(job/item['image'],destination/'previews'/f'slide-{i:02}.png')
    if spec.get('designSystem')=='modern-scientific-editorial':shutil.copyfile(ROOT/'design-system.md',destination/'design-system.md')
    for f in ['D/design-report.json','D/scientific-visualization-audit.json','E/native-check.json','E/build-manifest.json']:
        # Legacy 2.0 fixtures predate the palette audit. Editorial output must
        # retain it; do not make an absent legacy-only artifact break finalizing.
        if f=='D/scientific-visualization-audit.json' and spec.get('designSystem')!='modern-scientific-editorial' and not (job/f).is_file():continue
        target=destination/'audit'/f;target.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(job/f,target)
    for source in ['B/analysis.json','B/review.json','C/plan.json','C/review.json','F/qa-report.json','F/visual-review.json']:
        target=destination/'audit'/source
        target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copyfile(job/source,target)
    if spec.get('editorialQualityPolicy'):
        for source in ['D/editorial-quality.json','F/editorial-quality.json']:
            target=destination/'audit'/source;target.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(job/source,target)
    rebuild=read(job/'D/deck-spec.json')
    for slide in rebuild['slides']:
        for obj in slide['objects']:
            if obj['type']=='image':
                target=f'assets/{obj["assetId"]}{Path(obj["path"]).suffix}'
                (destination/'assets').mkdir(exist_ok=True)
                shutil.copyfile(obj['path'],destination/target)
                obj['path']=target
    write(destination/'deck-spec.json',rebuild)
    shutil.copyfile(ROOT/'scripts/studio_native.cjs',destination/'studio_native.cjs')
    (destination/'rebuild.cjs').write_text("const path=require('path');process.chdir(__dirname);require('./studio_native.cjs').build('deck-spec.json',path.resolve('rebuilt.pptx'));\n",encoding='utf-8')
    (destination/'visual-quality-report.md').write_text('# 视觉质量报告\n\n实际 PowerPoint 渲染逐页检查已记录。\n\n'+
        '\n'.join(f'- {s["id"]}: {s["status"]} — {s["notes"]}' for s in review['slides'])+
        '\n\n审阅者：'+review['reviewer']+'\n\n机器检查边界：\n'+ '\n'.join('- '+w for w in qa_report['warnings']),encoding='utf-8')
    write(destination/'manifest.json',{'purpose':read(job/'run.json')['settings']['purpose'],'sourceSha256':read(job/'run.json')['source']['sha256'],
        'scientificReviewSha256':sha(job/'B/review.json'),'narrativeReviewSha256':sha(job/'C/review.json'),
        'visualReviewSha256':sha(job/'F/visual-review.json'),
        'files':{str(f.relative_to(destination)).replace('\\','/'):sha(f) for f in destination.rglob('*') if f.is_file()}})
    return destination
